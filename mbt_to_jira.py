#!/usr/bin/env python3
"""MBT問合せExcelからJira起票を行うコンソールユーティリティ。"""

from __future__ import annotations

import base64
import json
import os
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any


SHEET_NAME = "MBTTool_問合せ"
DATA_START_ROW = 3
LOGIN_INFO_PATH = r"C:\問合せJira起票\loginInfo.txt"
PROJECT_KEY = "MBTG"
ISSUE_TYPE_STORY = "ストーリー"
ISSUE_TYPE_BUG = "バグ"

CATEGORY_MAP = {
    "Model Generator": "Modeling_Unit",
    "T/C Generator": "Generator_Unit",
}

BUG_FIELD_NAMES = {
    "symptom": "事象 (Hiện tượng)",
    "steps": "再現手順 (Các bước tái hiện)",
    "version": "不具合を検出したバージョン(version phát sinh bug)",
    "occurred_at": "発生日時(Thời điểm phát sinh)",
    "expected": "期待結果 (Kết quả mong đợi)",
}
CATEGORY_FIELD_NAME = "Category"


class AppError(Exception):
    pass


@dataclass
class JiraCredentials:
    base_url: str
    email: str
    api_token: str


@dataclass
class RowData:
    row_index: int
    management_no: str
    tool: str
    title: str
    detail: str
    env_info: str
    reproduction_steps: str
    expected_result: str
    ticket_no: str


def prompt_input(message: str) -> str | None:
    value = input(message).strip()
    if value == "" or value.lower() == "q":
        return None
    return value


def parse_login_info(path: str) -> JiraCredentials:
    file_path = Path(path)
    if not file_path.exists():
        raise AppError(f"認証情報ファイルが見つかりません: {path}")

    values: dict[str, str] = {}
    try:
        for raw_line in file_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" not in line:
                continue
            key, val = line.split("=", 1)
            values[key.strip()] = val.strip()
    except OSError as exc:
        raise AppError("認証情報ファイルの読み込みに失敗しました。") from exc

    missing = [k for k in ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN") if not values.get(k)]
    if missing:
        raise AppError(f"認証情報ファイルの必須キーが不足しています: {', '.join(missing)}")

    return JiraCredentials(
        base_url=values["JIRA_BASE_URL"].rstrip("/"),
        email=values["JIRA_EMAIL"],
        api_token=values["JIRA_API_TOKEN"],
    )


def get_excel_path() -> Path:
    path = os.getenv("MBT_INQUIRY_XLSX_PATH", "").strip()
    if not path:
        raise AppError("環境変数 MBT_INQUIRY_XLSX_PATH が未設定です。")

    file_path = Path(path)
    if not file_path.exists():
        raise AppError(f"Excelファイルが存在しません: {file_path}")
    return file_path


def _cell_str(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_target_row(excel_path: Path, management_no: str) -> tuple[Any, Any, RowData]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise AppError("openpyxl がインストールされていません。`pip install openpyxl` を実行してください。") from exc

    try:
        wb = load_workbook(excel_path)
    except Exception as exc:
        raise AppError(f"Excel読み込みに失敗しました: {excel_path}") from exc

    if SHEET_NAME not in wb.sheetnames:
        raise AppError(f"シートが見つかりません: {SHEET_NAME}")

    ws = wb[SHEET_NAME]
    matches: list[RowData] = []
    for row in range(DATA_START_ROW, ws.max_row + 1):
        if _cell_str(ws.cell(row=row, column=1).value) != management_no:
            continue
        matches.append(
            RowData(
                row_index=row,
                management_no=_cell_str(ws.cell(row=row, column=1).value),
                tool=_cell_str(ws.cell(row=row, column=5).value),
                title=_cell_str(ws.cell(row=row, column=6).value),
                detail=_cell_str(ws.cell(row=row, column=7).value),
                env_info=_cell_str(ws.cell(row=row, column=8).value),
                reproduction_steps=_cell_str(ws.cell(row=row, column=9).value),
                expected_result=_cell_str(ws.cell(row=row, column=10).value),
                ticket_no=_cell_str(ws.cell(row=row, column=19).value),
            )
        )

    if len(matches) == 0:
        raise AppError(f"管理番号 {management_no} は見つかりませんでした。")
    if len(matches) > 1:
        raise AppError(f"管理番号 {management_no} が複数行に存在します。")

    return wb, ws, matches[0]


def _auth_header(creds: JiraCredentials) -> str:
    raw = f"{creds.email}:{creds.api_token}".encode("utf-8")
    return "Basic " + base64.b64encode(raw).decode("ascii")


def request_json_with_retry(
    method: str,
    url: str,
    creds: JiraCredentials,
    payload: dict[str, Any] | None = None,
) -> tuple[int, dict[str, Any] | list[Any] | str]:
    wait_seconds = [1, 2, 4]
    headers = {
        "Authorization": _auth_header(creds),
        "Accept": "application/json",
    }

    body_bytes = None
    if payload is not None:
        body_bytes = json.dumps(payload).encode("utf-8")
        headers["Content-Type"] = "application/json"

    for attempt in range(1, 4):
        req = urllib.request.Request(url=url, data=body_bytes, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                code = resp.getcode()
                text = resp.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as exc:
            code = exc.code
            text = exc.read().decode("utf-8", errors="replace")
        except urllib.error.URLError as exc:
            if attempt == 3:
                raise AppError(f"Jira接続に失敗しました: {exc.reason}") from exc
            time.sleep(wait_seconds[attempt - 1])
            continue

        if code in (401, 403):
            raise AppError("Jira認証エラーが発生しました（401/403）。loginInfo.txt を確認してください。")

        if code == 429 or 500 <= code < 600:
            if attempt == 3:
                raise AppError(f"Jira API一時エラー: HTTP {code} {text[:200]}")
            time.sleep(wait_seconds[attempt - 1])
            continue

        if text:
            try:
                return code, json.loads(text)
            except json.JSONDecodeError:
                return code, text
        return code, {}

    raise AppError("Jira API呼び出しに失敗しました。")


def resolve_custom_fields(creds: JiraCredentials, issue_type: str) -> dict[str, str]:
    code, body = request_json_with_retry("GET", f"{creds.base_url}/rest/api/3/field", creds)
    if code >= 400:
        raise AppError(f"Jiraフィールド一覧の取得に失敗しました: HTTP {code}")
    if not isinstance(body, list):
        raise AppError("Jiraフィールド一覧の形式が不正です。")

    by_name: dict[str, str] = {}
    for item in body:
        if isinstance(item, dict) and item.get("name") and item.get("id"):
            by_name[str(item["name"])] = str(item["id"])

    required = {"category": CATEGORY_FIELD_NAME}
    if issue_type == ISSUE_TYPE_BUG:
        required.update(BUG_FIELD_NAMES)

    resolved: dict[str, str] = {}
    for key, display_name in required.items():
        field_id = by_name.get(display_name)
        if not field_id:
            raise AppError(f"Jiraフィールド名を解決できませんでした: {display_name}")
        resolved[key] = field_id
    return resolved


def build_issue_fields(row: RowData, issue_type: str, custom_fields: dict[str, str]) -> dict[str, Any]:
    fields: dict[str, Any] = {
        "project": {"key": PROJECT_KEY},
        "issuetype": {"name": issue_type},
        "summary": row.title,
        "labels": [row.management_no],
    }

    category_value = CATEGORY_MAP.get(row.tool)
    if category_value:
        fields[custom_fields["category"]] = {"value": category_value}
    elif row.tool:
        print(f"[WARN] 未定義の対象ツールのためCategoryは設定しません: {row.tool}")

    if issue_type == ISSUE_TYPE_STORY:
        if row.detail:
            fields["description"] = {
                "type": "doc",
                "version": 1,
                "content": [
                    {
                        "type": "paragraph",
                        "content": [
                            {
                                "type": "text",
                                "text": row.detail,
                            }
                        ],
                    }
                ],
            }
    elif issue_type == ISSUE_TYPE_BUG:
        if row.detail:
            fields[custom_fields["symptom"]] = row.detail
        if row.reproduction_steps:
            fields[custom_fields["steps"]] = row.reproduction_steps
        if row.env_info:
            fields[custom_fields["version"]] = row.env_info
            fields[custom_fields["occurred_at"]] = row.env_info
        if row.expected_result:
            fields[custom_fields["expected"]] = row.expected_result
    else:
        raise AppError(f"未対応のIssueTypeです: {issue_type}")

    return fields


def print_preview(fields: dict[str, Any]) -> None:
    print("--- Dry-run Preview ---")
    for key, value in fields.items():
        print(f"{key}: {value}")
    print("-----------------------")


def create_issue(creds: JiraCredentials, fields: dict[str, Any]) -> str:
    code, body = request_json_with_retry("POST", f"{creds.base_url}/rest/api/3/issue", creds, {"fields": fields})
    if code >= 400:
        raise AppError(f"Jira起票に失敗しました: HTTP {code}")
    if not isinstance(body, dict) or not body.get("key"):
        raise AppError("Jira起票レスポンスに課題キーが含まれていません。")
    return str(body["key"])


def write_back_ticket_no(workbook: Any, worksheet: Any, row_index: int, ticket_key: str, excel_path: Path) -> None:
    worksheet.cell(row=row_index, column=19, value=ticket_key)
    try:
        workbook.save(excel_path)
    except Exception as exc:
        raise AppError(f"Excel保存に失敗しました: {excel_path}") from exc


def main() -> int:
    try:
        management_no = prompt_input("管理番号を入力してください（qで終了）: ")
        if management_no is None:
            print("終了します。")
            return 0

        issue_type = prompt_input("IssueTypeを入力してください（ストーリー/バグ, qで終了）: ")
        if issue_type is None:
            print("終了します。")
            return 0
        if issue_type not in (ISSUE_TYPE_STORY, ISSUE_TYPE_BUG):
            raise AppError("IssueTypeは ストーリー または バグ を指定してください。")

        dry_run_input = prompt_input("Dry-runですか？（y/n, qで終了）: ")
        if dry_run_input is None:
            print("終了します。")
            return 0
        if dry_run_input.lower() not in ("y", "n"):
            raise AppError("Dry-runは y または n を指定してください。")
        dry_run = dry_run_input.lower() == "y"

        print(f"[INFO] 管理番号={management_no}, IssueType={issue_type}, dry-run={dry_run}")

        excel_path = get_excel_path()
        creds = parse_login_info(LOGIN_INFO_PATH)
        workbook, worksheet, row = load_target_row(excel_path, management_no)

        if row.ticket_no:
            print(f"[INFO] 起票済みのためスキップします。S列={row.ticket_no}")
            print("[RESULT] SKIPPED")
            return 0

        if not row.title:
            raise AppError(f"管理番号 {row.management_no} のタイトル(F列)が空のため起票できません。")

        custom_fields = resolve_custom_fields(creds, issue_type)
        fields = build_issue_fields(row, issue_type, custom_fields)

        if dry_run:
            print_preview(fields)
            print("[RESULT] DRY-RUN")
            return 0

        issue_key = create_issue(creds, fields)
        print(f"[INFO] Jira起票成功: {issue_key}")
        write_back_ticket_no(workbook, worksheet, row.row_index, issue_key, excel_path)
        print("[RESULT] SUCCESS")
        return 0

    except AppError as exc:
        print(f"[ERROR] {exc}")
        print("[RESULT] FAILED")
        return 1
    except KeyboardInterrupt:
        print("\n中断しました。")
        return 1


if __name__ == "__main__":
    sys.exit(main())
